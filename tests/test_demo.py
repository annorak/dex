import hashlib
import json
import shutil
import sqlite3
import subprocess
import sys
from collections.abc import Iterator
from pathlib import Path
from typing import Any
from xml.etree import ElementTree
from zipfile import ZIP_STORED, ZipFile

import pytest
import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "demo"
BASELINE = ROOT / "tests/demo.sha256"

type ExpectedRow = tuple[str | int | float, ...]

DATABASE_TABLES: list[tuple[str, str, tuple[str, ...], list[ExpectedRow]]] = [
    (
        "cerulean/accounting.db",
        "invoices",
        (
            "invoice_id",
            "community_id",
            "vendor_id",
            "invoice_date",
            "total_cents",
            "status",
        ),
        [
            ("I099", "H17", "V7", "2025-11-03", 15000, "C"),
            ("I100", "H17", "V7", "2026-07-10", 125000, "C"),
            ("I101", "H17", "V7", "2026-08-01", 20000, "C"),
            ("I102", "H17", "V7", "2026-08-15", 10000, "V"),
            ("I103", "H17", "V3", "2026-06-01", 40000, "C"),
            ("I200", "H71", "V7", "2026-08-20", 99900, "C"),
        ],
    ),
    (
        "cerulean/accounting.db",
        "invoice_lines",
        ("invoice_id", "line_no", "description", "amount_cents"),
        [
            ("I099", 1, "Invoice total", 15000),
            ("I100", 1, "LED fixtures", 100000),
            ("I100", 2, "Installation labor", 25000),
            ("I101", 1, "Invoice total", 20000),
            ("I102", 1, "Invoice total", 10000),
            ("I103", 1, "Invoice total", 40000),
            ("I200", 1, "Invoice total", 99900),
        ],
    ),
    (
        "cerulean/accounting.db",
        "payments",
        ("payment_id", "invoice_id", "paid_date", "amount_cents"),
        [
            ("PAY1", "I099", "2025-12-01", 15000),
            ("PAY2", "I103", "2026-06-20", 40000),
            ("PAY3", "I100", "2026-07-25", 20000),
            ("PAY4", "I100", "2026-09-05", 30000),
            ("PAY5", "I101", "2026-09-12", 20000),
        ],
    ),
    (
        "vermilion/bookings.db",
        "reservations",
        ("reservation_id", "account_id", "traveler_id", "depart_date", "status"),
        [
            ("R100", "A01", "T07", "2026-10-20", "C"),
            ("R101", "A01", "T07", "2026-10-22", "X"),
            ("R102", "A01", "T07", "2026-09-25", "C"),
            ("R103", "A01", "T07", "2026-10-08", "C"),
            ("R104", "A01", "T07", "2026-11-03", "X"),
            ("R200", "A02", "T07", "2026-10-20", "C"),
            ("R201", "A02", "T07", "2026-10-15", "X"),
            ("R300", "A01", "T07", "2026-10-22", "P"),
        ],
    ),
    (
        "pallet/accounting.db",
        "communities",
        ("community_id", "name"),
        [
            ("P01", "Pallet Grove"),
            ("P02", "Lavender Grove"),
        ],
    ),
    (
        "pallet/accounting.db",
        "vendors",
        ("vendor_id", "name"),
        [("V1", "Eevee Maintenance")],
    ),
    (
        "pallet/accounting.db",
        "invoices",
        (
            "invoice_id",
            "community_id",
            "vendor_id",
            "invoice_date",
            "total_cents",
            "balance_cents",
            "status",
        ),
        [
            ("P100", "P01", "V1", "2026-09-01", 20000, 15000, "O"),
            ("P101", "P01", "V1", "2026-08-01", 10000, 0, "C"),
        ],
    ),
]
WORKBOOK_TABLES: list[tuple[str, str, tuple[str, ...], list[ExpectedRow]]] = [
    (
        "cerulean/directory.xlsx",
        "Communities",
        ("community_id", "name"),
        [
            ("H17", "Viridian Court"),
            ("H71", "Viridian Court II"),
            ("H18", "Pewter Court"),
        ],
    ),
    (
        "cerulean/directory.xlsx",
        "Vendors",
        ("community_id", "vendor_id", "name"),
        [
            ("H17", "V3", "Bulbasaur Landscaping"),
            ("H17", "V7", "Pikachu Lighting"),
            ("H71", "V7", "Squirtle Pool Care"),
        ],
    ),
    (
        "cerulean/Open Invoices - Aug 31 2026.xlsx",
        "Open Invoices",
        ("Community", "Vendor", "Invoice", "Amount Open"),
        [
            ("Viridian Court", "Pikachu Lighting", "I100", 1050.00),
            ("Viridian Court", "Pikachu Lighting", "I101", 200.00),
            ("Viridian Court II", "Squirtle Pool Care", "I200", 999.00),
        ],
    ),
    (
        "vermilion/clients.xlsx",
        "Accounts",
        ("account_id", "name"),
        [("A01", "Silph Account"), ("A02", "Devon Account")],
    ),
    (
        "vermilion/clients.xlsx",
        "Travelers",
        ("account_id", "traveler_id", "name"),
        [
            ("A01", "T07", "Ash Ketchum"),
            ("A01", "T09", "Brock"),
            ("A02", "T07", "Ash Ketchum"),
        ],
    ),
    (
        "vermilion/Trips (old).xlsx",
        "Trips",
        ("Traveler", "Account", "Reservation", "Depart", "Status"),
        [
            ("Ash Ketchum", "Silph Account", "R100", "2026-10-20", "Confirmed"),
            ("Ash Ketchum", "Silph Account", "R101", "2026-10-22", "Confirmed"),
        ],
    ),
]
EXPECTED_DOCS = {
    "cerulean/docs/ap_process.md": (
        "1. Vendor bills are entered as invoices, and each payment "
        "is recorded as its own entry against an invoice.\n2. All "
        "amounts in the accounting system are in cents.\n3. Each "
        'month, the AP clerk exports an "Open Invoices" sheet for '
        "the board.\n"
    ),
    "vermilion/docs/booking_process.md": (
        "1. Cancelled bookings are marked `X`.\n2. Each client "
        "company has its own account, and traveler numbers start "
        "over in each account.\n"
    ),
    "pallet/docs/ap_process.md": (
        "1. Invoices store the amount still owed in balance_cents. "
        "Amounts are whole US cents.\n2. Invoice status O means "
        "open. C means closed and paid in full.\n3. Vendor IDs are "
        "unique across Pallet HOA. The accounting database is the "
        "trusted source.\n"
    ),
}
EXPECTED_EMPLOYEES = {
    "cerulean": [
        (
            "status_meaning",
            "accounting.invoices.status",
            (
                "C means current, so the invoice is posted and counts. V "
                "means void, so ignore it."
            ),
        ),
        (
            "balance_rule",
            "accounting.invoices",
            (
                "There's no balance field. What's owed is the total minus "
                "the payments recorded against the invoice. An invoice is "
                "unpaid if it isn't void and something is still owed."
            ),
        ),
        (
            "trusted_source",
            "open_invoices_aug",
            (
                "Use the accounting system for anything about balances. The"
                " Open Invoices sheet is a monthly export for the board, so"
                " it goes stale."
            ),
        ),
        (
            "id_scope",
            "directory.vendors.vendor_id",
            (
                "Vendor numbers are assigned per community. V7 at one "
                "community is a different vendor from V7 at another."
            ),
        ),
        (
            "current_process",
            "unpaid_invoices_by_community",
            (
                "Board members ask before their meetings. The AP clerk "
                "emails them the latest Open Invoices sheet."
            ),
        ),
    ],
    "vermilion": [
        (
            "status_meaning",
            "bookings.reservations.status",
            (
                "C means confirmed (booked). P means pending: on hold, not "
                "booked yet. X means cancelled."
            ),
        ),
        (
            "trusted_source",
            "trips_old",
            (
                "Use the bookings system. The Trips sheet stopped being "
                "updated in September."
            ),
        ),
        (
            "date_rule",
            "bookings.reservations.depart_date",
            "It's the day the trip starts. We don't track times here.",
        ),
        (
            "current_process",
            "upcoming_reservations",
            "Account managers check the bookings system before a client call.",
        ),
    ],
    "pallet": [
        (
            "status_meaning",
            "accounting.invoices.status",
            "O means open. C means closed and paid in full.",
        ),
        (
            "balance_rule",
            "accounting.invoices.balance_cents",
            (
                "Use the stored balance_cents. A positive balance on an "
                "open invoice is unpaid."
            ),
        ),
        (
            "id_scope",
            "accounting.vendors.vendor_id",
            "Vendor IDs are unique across Pallet HOA.",
        ),
        ("trusted_source", "accounting", "Use the accounting database for balances."),
    ],
}


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def file_hashes(folder: Path) -> dict[str, str]:
    return {
        path.relative_to(folder).as_posix(): hashlib.sha256(
            path.read_bytes()
        ).hexdigest()
        for path in sorted(folder.rglob("*"))
        if path.is_file()
    }


def read_baseline(path: Path) -> dict[str, str]:
    return {
        name: digest
        for digest, name in (
            line.split("  ", 1)
            for line in path.read_text(encoding="utf-8").splitlines()
        )
    }


def assert_hash_baseline(folder: Path, baseline: Path) -> None:
    assert file_hashes(folder) == read_baseline(baseline), "Hash baseline mismatch"


@pytest.fixture(scope="module", params=["fresh", "checked_in"])
def generated(
    tmp_path_factory: pytest.TempPathFactory, request: pytest.FixtureRequest
) -> Path:
    if request.param == "checked_in":
        return DEMO
    folder = tmp_path_factory.mktemp("generated") / "demo"
    subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts/generate_demo.py"),
            "--output",
            str(folder),
        ],
        cwd=ROOT,
        check=True,
        timeout=30,
    )
    return folder


@pytest.fixture
def source_db() -> Iterator[sqlite3.Connection]:
    connection = sqlite3.connect(
        f"{(DEMO / 'companies/cerulean/accounting.db').as_uri()}?mode=ro", uri=True
    )
    try:
        yield connection
    finally:
        connection.close()


@pytest.mark.parametrize(
    ("file", "table", "columns", "rows"),
    DATABASE_TABLES,
    ids=[f"{file}:{table}:{len(rows)}rows" for file, table, _, rows in DATABASE_TABLES],
)
def test_every_database_row(
    generated: Path,
    file: str,
    table: str,
    columns: tuple[str, ...],
    rows: list[ExpectedRow],
) -> None:
    connection = sqlite3.connect(
        f"{(generated / 'companies' / file).as_uri()}?mode=ro", uri=True
    )
    try:
        assert connection.execute("PRAGMA integrity_check").fetchall() == [("ok",)]
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
        cursor = connection.execute(f'SELECT * FROM "{table}" ORDER BY rowid')
        assert tuple(item[0] for item in cursor.description) == columns
        assert cursor.fetchall() == rows
    finally:
        connection.close()


@pytest.mark.parametrize(
    ("file", "tab", "columns", "rows"),
    WORKBOOK_TABLES,
    ids=[f"{file}:{tab}:{len(rows)}rows" for file, tab, _, rows in WORKBOOK_TABLES],
)
def test_every_workbook_row(
    generated: Path,
    file: str,
    tab: str,
    columns: tuple[str, ...],
    rows: list[ExpectedRow],
) -> None:
    workbook = load_workbook(
        generated / "companies" / file, read_only=True, data_only=True
    )
    try:
        actual = list(workbook[tab].values)
        assert actual[0] == columns
        assert actual[1:] == rows
        assert all(cell is not None for row in actual for cell in row)
    finally:
        workbook.close()


@pytest.mark.parametrize("file", EXPECTED_DOCS)
def test_source_docs_only_disclose_documented_facts(generated: Path, file: str) -> None:
    assert (generated / "companies" / file).read_text(
        encoding="utf-8"
    ) == EXPECTED_DOCS[file]


@pytest.mark.parametrize("company", EXPECTED_EMPLOYEES)
def test_exact_employee_answers(generated: Path, company: str) -> None:
    actual = yaml.safe_load(
        (generated / f"keys/{company}.employee.yaml").read_text(encoding="utf-8")
    )
    assert actual == {
        "company": company,
        "answers": [
            {"topic": topic, "target": target, "answer": answer}
            for topic, target, answer in EXPECTED_EMPLOYEES[company]
        ],
    }


def test_exact_hoa_v1(generated: Path) -> None:
    actual = yaml.safe_load(
        (generated / "checklists/hoa/unpaid_invoices.v1.yaml").read_text(
            encoding="utf-8"
        )
    )
    assert actual == {
        "industry": "hoa",
        "lookup": "unpaid_invoices_by_community",
        "version": 1,
        "question": (
            "Which invoices are still unpaid for a community, and which"
            " vendors are they for?"
        ),
        "inputs": {"community_name": "the exact name of one community"},
        "outputs": [
            "invoice_id",
            "vendor_id",
            "vendor_name",
            "outstanding_cents",
            "currency",
        ],
        "ask_employee": [
            "What does each invoice status code mean? Codes differ between companies.",
            "Is a balance stored on each invoice, or is it worked out from payments?",
            "If two sources disagree about what's owed, which one do you trust?",
        ],
        "watch_for": [
            (
                "Community names can look alike. Match the exact name, then"
                " confirm the ID."
            ),
            "Check whether amounts are stored in cents or dollars.",
        ],
        "checks": ["identity_is_unique", "exact_rows_and_values", "known_empty_is_ok"],
    }
    assert list((generated / "checklists").rglob("*.yaml")) == [
        generated / "checklists/hoa/unpaid_invoices.v1.yaml"
    ]


@pytest.mark.parametrize(
    ("company", "date", "sources"),
    [
        (
            "cerulean",
            "2026-09-30",
            [
                {
                    "id": "accounting",
                    "file": "accounting.db",
                    "tables": ["invoices", "invoice_lines", "payments"],
                },
                {
                    "id": "directory",
                    "file": "directory.xlsx",
                    "tabs": ["Communities", "Vendors"],
                },
                {
                    "id": "open_invoices_aug",
                    "file": "Open Invoices - Aug 31 2026.xlsx",
                    "tabs": ["Open Invoices"],
                },
                {"id": "docs", "file": "docs/ap_process.md"},
            ],
        ),
        (
            "vermilion",
            "2026-10-08",
            [
                {"id": "bookings", "file": "bookings.db", "tables": ["reservations"]},
                {
                    "id": "clients",
                    "file": "clients.xlsx",
                    "tabs": ["Accounts", "Travelers"],
                },
                {"id": "trips_old", "file": "Trips (old).xlsx", "tabs": ["Trips"]},
                {"id": "docs", "file": "docs/booking_process.md"},
            ],
        ),
        (
            "pallet",
            "2026-09-30",
            [
                {
                    "id": "accounting",
                    "file": "accounting.db",
                    "tables": ["communities", "vendors", "invoices"],
                },
                {"id": "docs", "file": "docs/ap_process.md"},
            ],
        ),
    ],
)
def test_source_inventory(
    generated: Path, company: str, date: str, sources: list[dict[str, Any]]
) -> None:
    folder = generated / "companies" / company
    assert read_json(folder / "sources.json") == {
        "company": company,
        "snapshot": f"{company}-v1",
        "as_of": date,
        "sources": sources,
    }
    assert {
        path.relative_to(folder).as_posix()
        for path in folder.rglob("*")
        if path.is_file()
    } == {
        "sources.json",
        *[source["file"] for source in sources],
    }
    for source in sources:
        path = folder / source["file"]
        if "tables" in source:
            connection = sqlite3.connect(f"{path.as_uri()}?mode=ro", uri=True)
            try:
                assert sorted(
                    row[0]
                    for row in connection.execute(
                        "SELECT name FROM sqlite_master WHERE type='table'"
                    )
                ) == sorted(source["tables"])
            finally:
                connection.close()
        if "tabs" in source:
            workbook = load_workbook(path, read_only=True)
            try:
                assert workbook.sheetnames == source["tabs"]
            finally:
                workbook.close()


def assert_cerulean_keys(path: Path) -> None:
    keys = read_json(path)
    main_row = {
        "invoice_id": "cerulean/H17/I100",
        "vendor_id": "cerulean/H17/V7",
        "vendor_name": "Pikachu Lighting",
        "outstanding_cents": 75000,
        "currency": "USD",
        "source_refs": [
            {
                "source": "accounting",
                "table": "invoices",
                "record_id": "cerulean/H17/I100",
            },
            {
                "source": "accounting",
                "table": "payments",
                "record_id": "cerulean/H17/PAY3",
            },
            {
                "source": "accounting",
                "table": "payments",
                "record_id": "cerulean/H17/PAY4",
            },
        ],
    }
    other_row = {
        "invoice_id": "cerulean/H71/I200",
        "vendor_id": "cerulean/H71/V7",
        "vendor_name": "Squirtle Pool Care",
        "outstanding_cents": 99900,
        "currency": "USD",
        "source_refs": [
            {
                "source": "accounting",
                "table": "invoices",
                "record_id": "cerulean/H71/I200",
            }
        ],
    }
    expected_cases: list[tuple[str, str, str, list[str], list[dict[str, Any]]]] = [
        ("identity_viridian", "Viridian Court", "ok", ["cerulean/H17"], [main_row]),
        ("rows_viridian", "Viridian Court", "ok", ["cerulean/H17"], [main_row]),
        ("values_viridian", "Viridian Court", "ok", ["cerulean/H17"], [main_row]),
        (
            "lookalike_viridian_ii",
            "Viridian Court II",
            "ok",
            ["cerulean/H71"],
            [other_row],
        ),
        ("empty_pewter", "Pewter Court", "ok", ["cerulean/H18"], []),
        ("unknown_name", "Viridian Ct", "unresolved_identity", [], []),
        ("source_refs", "Viridian Court", "ok", ["cerulean/H17"], [main_row]),
    ]
    assert keys["company"] == "cerulean"
    assert (keys["snapshot"], keys["as_of"], keys["lookup"]) == (
        "cerulean-v1",
        "2026-09-30",
        "unpaid_invoices_by_community",
    )
    assert keys["checks"] == [
        {
            "check": check,
            "input": {"community_name": name},
            "status": status,
            "identity": identity,
            "rows": rows,
        }
        for check, name, status, identity, rows in expected_cases
    ]
    assert keys["main"] == keys["checks"][1]
    assert keys["follow_up"] == {
        "question": (
            "How much has each Viridian Court vendor been paid in 2026,"
            " counting payments dated January 1 through September 30?"
        ),
        "input": {
            "community_name": "Viridian Court",
            "from_date": "2026-01-01",
            "through_date": "2026-09-30",
        },
        "status": "ok",
        "rows": [
            {
                "vendor_id": "cerulean/H17/V3",
                "vendor_name": "Bulbasaur Landscaping",
                "paid_cents": 40000,
                "currency": "USD",
            },
            {
                "vendor_id": "cerulean/H17/V7",
                "vendor_name": "Pikachu Lighting",
                "paid_cents": 70000,
                "currency": "USD",
            },
        ],
    }


def test_cerulean_keys(generated: Path) -> None:
    assert_cerulean_keys(generated / "keys/cerulean.answers.json")


def test_vermilion_keys(generated: Path) -> None:
    keys = read_json(generated / "keys/vermilion.answers.json")
    main_row = {
        "reservation_id": "vermilion/A01/R100",
        "traveler_id": "vermilion/A01/T07",
        "traveler_name": "Ash Ketchum",
        "account_id": "vermilion/A01",
        "account_name": "Silph Account",
        "depart_date": "2026-10-20",
        "source_refs": [
            {
                "source": "bookings",
                "table": "reservations",
                "record_id": "vermilion/A01/R100",
            }
        ],
    }
    other_row = {
        "reservation_id": "vermilion/A02/R200",
        "traveler_id": "vermilion/A02/T07",
        "traveler_name": "Ash Ketchum",
        "account_id": "vermilion/A02",
        "account_name": "Devon Account",
        "depart_date": "2026-10-20",
        "source_refs": [
            {
                "source": "bookings",
                "table": "reservations",
                "record_id": "vermilion/A02/R200",
            }
        ],
    }
    expected_cases: list[tuple[str, str, str, str, list[str], list[dict[str, Any]]]] = [
        (
            "identity_ash_silph",
            "Ash Ketchum",
            "Silph Account",
            "ok",
            ["vermilion/A01/T07"],
            [main_row],
        ),
        (
            "rows_ash_silph",
            "Ash Ketchum",
            "Silph Account",
            "ok",
            ["vermilion/A01/T07"],
            [main_row],
        ),
        (
            "values_ash_silph",
            "Ash Ketchum",
            "Silph Account",
            "ok",
            ["vermilion/A01/T07"],
            [main_row],
        ),
        (
            "same_name_devon",
            "Ash Ketchum",
            "Devon Account",
            "ok",
            ["vermilion/A02/T07"],
            [other_row],
        ),
        ("empty_brock", "Brock", "Silph Account", "ok", ["vermilion/A01/T09"], []),
        ("unknown_traveler", "Ash K.", "Silph Account", "unresolved_identity", [], []),
        (
            "source_refs",
            "Ash Ketchum",
            "Silph Account",
            "ok",
            ["vermilion/A01/T07"],
            [main_row],
        ),
    ]
    assert keys["company"] == "vermilion"
    assert (keys["snapshot"], keys["as_of"], keys["lookup"]) == (
        "vermilion-v1",
        "2026-10-08",
        "upcoming_reservations",
    )
    assert keys["checks"] == [
        {
            "check": check,
            "input": {
                "traveler_name": traveler,
                "account_name": account,
                "after_date": "2026-10-08",
            },
            "status": status,
            "identity": identity,
            "rows": rows,
        }
        for check, traveler, account, status, identity, rows in expected_cases
    ]
    assert keys["main"] == keys["checks"][1]
    assert keys["follow_up"] == {
        "question": (
            "Which Silph Account travelers have a cancelled reservation"
            " that was scheduled to depart in October 2026?"
        ),
        "input": {
            "account_name": "Silph Account",
            "from_date": "2026-10-01",
            "through_date": "2026-10-31",
        },
        "status": "ok",
        "rows": [
            {
                "traveler_id": "vermilion/A01/T07",
                "traveler_name": "Ash Ketchum",
                "account_id": "vermilion/A01",
                "account_name": "Silph Account",
            }
        ],
    }


def test_pallet_keys(generated: Path) -> None:
    keys = read_json(generated / "keys/pallet.answers.json")
    assert (keys["company"], keys["snapshot"], keys["as_of"]) == (
        "pallet",
        "pallet-v1",
        "2026-09-30",
    )
    assert keys["lookup"] == "unpaid_invoices_by_community"
    assert keys["main"] == {
        "check": "rows_pallet",
        "input": {"community_name": "Pallet Grove"},
        "status": "ok",
        "identity": ["pallet/P01"],
        "rows": [
            {
                "invoice_id": "pallet/P01/P100",
                "vendor_id": "pallet/P01/V1",
                "vendor_name": "Eevee Maintenance",
                "outstanding_cents": 15000,
                "currency": "USD",
                "source_refs": [
                    {
                        "source": "accounting",
                        "table": "invoices",
                        "record_id": "pallet/P01/P100",
                    }
                ],
            }
        ],
    }
    assert keys["empty"] == {
        "check": "empty_lavender",
        "input": {"community_name": "Lavender Grove"},
        "status": "ok",
        "identity": ["pallet/P02"],
        "rows": [],
    }
    assert keys["unknown"] == {
        "check": "unknown_name",
        "input": {"community_name": "Pallet Gr"},
        "status": "unresolved_identity",
        "identity": [],
        "rows": [],
    }


def test_double_count_table_and_fixture_traps(source_db: sqlite3.Connection) -> None:
    actual = source_db.execute(
        "SELECT l.line_no, l.amount_cents, p.payment_id, "
        "p.amount_cents FROM invoice_lines l JOIN payments p ON "
        "p.invoice_id=l.invoice_id WHERE l.invoice_id='I100' ORDER "
        "BY l.line_no,p.payment_id"
    ).fetchall()
    assert actual == [
        (1, 100000, "PAY3", 20000),
        (1, 100000, "PAY4", 30000),
        (2, 25000, "PAY3", 20000),
        (2, 25000, "PAY4", 30000),
    ]
    assert 125000 - sum(row[3] for row in actual) == 25000
    assert source_db.execute(
        "SELECT total_cents-(SELECT SUM(amount_cents) FROM payments"
        " WHERE invoice_id='I100') FROM invoices WHERE "
        "invoice_id='I100'"
    ).fetchone() == (75000,)
    assert source_db.execute(
        "SELECT COUNT(*) FROM invoices WHERE invoice_id='I100' AND status='O'"
    ).fetchone() == (0,)


@pytest.mark.parametrize(
    "file",
    [
        "companies/cerulean/directory.xlsx",
        "companies/cerulean/Open Invoices - Aug 31 2026.xlsx",
        "companies/vermilion/clients.xlsx",
        "companies/vermilion/Trips (old).xlsx",
    ],
)
def test_workbook_time_and_zip_are_fixed(generated: Path, file: str) -> None:
    with ZipFile(generated / file) as archive:
        assert archive.namelist() == sorted(archive.namelist())
        assert all(
            member.date_time == (1980, 1, 1, 0, 0, 0)
            and member.compress_type == ZIP_STORED
            for member in archive.infolist()
        )
        root = ElementTree.fromstring(archive.read("docProps/core.xml"))
        for name in ("created", "modified"):
            value = root.find(f"{{http://purl.org/dc/terms/}}{name}")
            assert value is not None and value.text == "2026-09-30T00:00:00Z"
    if "Open Invoices" in file:
        workbook = load_workbook(generated / file)
        try:
            assert [
                workbook.worksheets[0].cell(row=row, column=4).number_format
                for row in range(2, 5)
            ] == ["0.00"] * 3
        finally:
            workbook.close()


def test_pallet_reviewed_package(generated: Path) -> None:
    package = generated / "packages/pallet"
    expected_files = {
        "manifest.json",
        "overview.md",
        "sources.json",
        "lookups.json",
        "decisions.json",
        "queries/identity.sql",
        "queries/answer.sql",
        "queries/settings.json",
        "checks/definitions.json",
        "checks/results.json",
        "evidence/accounting.db",
        "evidence/ap_process.md",
    }
    assert set(file_hashes(package)) == expected_files
    manifest = read_json(package / "manifest.json")
    assert manifest["files"] == {
        name: digest
        for name, digest in file_hashes(package).items()
        if name != "manifest.json"
    }
    assert (
        manifest["company"],
        manifest["snapshot"],
        manifest["as_of"],
        manifest["package_format"],
        manifest["dex_version"],
    ) == ("pallet", "pallet-v1", "2026-09-30", 1, "0.1.0")
    assert manifest["prepared_example"] is True and manifest["simulated"] is True
    assert manifest["checklist_used"] == {
        "industry": "hoa",
        "lookup": "unpaid_invoices_by_community",
        "version": 1,
    }
    assert manifest["query_version"] == manifest["checks_version"] == 1
    assert manifest["reviews"] == ["pallet-employee-1", "pallet-engineer-1"]
    assert manifest["limits"] == [
        "Fixed invented snapshot",
        "No live permissions or freshness validation",
        "No historical balances",
    ]
    for original, evidence in (
        ("accounting.db", "accounting.db"),
        ("docs/ap_process.md", "ap_process.md"),
    ):
        assert (package / "evidence" / evidence).read_bytes() == (
            generated / "companies/pallet" / original
        ).read_bytes()
    decisions = read_json(package / "decisions.json")["decisions"]
    assert [decision["id"] for decision in decisions] == manifest["reviews"]
    assert [decision["role"] for decision in decisions] == ["employee", "engineer"]
    assert all(
        decision["simulated"] is True and decision["outcome"] == "approved"
        for decision in decisions
    )
    assert all(decision["at"] == "2026-09-30T00:00:00Z" for decision in decisions)
    settings = read_json(package / "queries/settings.json")
    assert settings == {
        "database": "evidence/accounting.db",
        "read_only": True,
        "timeout_seconds": 2,
    }
    identity_sql = (package / "queries/identity.sql").read_text(encoding="utf-8")
    answer_sql = (package / "queries/answer.sql").read_text(encoding="utf-8")
    definitions = read_json(package / "checks/definitions.json")["checks"]
    results = read_json(package / "checks/results.json")
    assert results["run"] == "pallet-reviewed-v1"
    assert [check["id"] for check in definitions] == [
        "identity_is_unique",
        "exact_rows_and_values",
        "known_empty_is_ok",
    ]
    assert len(results["results"]) == 3
    connection = sqlite3.connect(
        f"{(package / settings['database']).as_uri()}?mode=ro", uri=True
    )
    try:
        for check, result in zip(definitions, results["results"], strict=True):
            expected_identity = (
                ["P02"] if check["id"] == "known_empty_is_ok" else ["P01"]
            )
            expected_rows = (
                []
                if check["id"] == "known_empty_is_ok"
                else [["P100", "V1", "Eevee Maintenance", 15000, "USD"]]
            )
            identity = [
                row[0]
                for row in connection.execute(
                    identity_sql, {"community_name": check["community_name"]}
                )
            ]
            actual = [
                list(row)
                for row in connection.execute(answer_sql, {"community_id": identity[0]})
            ]
            assert identity == check["expected_identity"] == expected_identity
            assert actual == check["expected_rows"] == expected_rows
            assert result == {
                "id": check["id"],
                "status": "passed",
                "actual_identity": expected_identity,
                "actual_rows": expected_rows,
            }
    finally:
        connection.close()
    sources = read_json(package / "sources.json")
    assert (sources["company"], sources["snapshot"], sources["as_of"]) == (
        "pallet",
        "pallet-v1",
        "2026-09-30",
    )
    assert sources["sources"][0] == {
        "id": "accounting",
        "original_file": "accounting.db",
        "file": "evidence/accounting.db",
        "tables": {
            "communities": {
                "primary_key": ["community_id"],
                "records": ["pallet/P01", "pallet/P02"],
            },
            "vendors": {"primary_key": ["vendor_id"], "records": ["pallet/P01/V1"]},
            "invoices": {
                "primary_key": ["invoice_id"],
                "records": ["pallet/P01/P100", "pallet/P01/P101"],
            },
        },
    }
    assert sources["sources"][1] == {
        "id": "docs",
        "original_file": "docs/ap_process.md",
        "file": "evidence/ap_process.md",
        "items": [1, 2, 3],
    }
    lookup = read_json(package / "lookups.json")["lookups"]
    assert len(lookup) == 1
    assert lookup[0]["id"] == "unpaid_invoices_by_community"
    assert lookup[0]["inputs"] == {"community_name": "the exact name of one community"}
    assert lookup[0]["outputs"] == [
        "invoice_id",
        "vendor_id",
        "vendor_name",
        "outstanding_cents",
        "currency",
    ]
    assert (
        lookup[0]["identity_query"] == "queries/identity.sql"
        and lookup[0]["answer_query"] == "queries/answer.sql"
    )
    assert lookup[0]["rules"] == [
        "Exact community name must resolve once.",
        "O is open; C is closed.",
        "Positive stored balance_cents on an open invoice is unpaid.",
        "Vendor IDs are unique across the company.",
        "Amounts are whole US cents; currency is USD.",
    ]
    overview = (package / "overview.md").read_text(encoding="utf-8")
    for fact in (
        "simulated review",
        "balance_cents",
        "O means open",
        "C means closed",
        "company-wide unique",
        "pallet-engineer-1",
        "three passing checks",
    ):
        assert fact in overview


def test_reference_sql_stays_hidden_and_is_executable(generated: Path) -> None:
    for company, file, answer_parameters, expected in (
        (
            "cerulean",
            "accounting.db",
            {"community_id": "H17"},
            [("I100", "V7", "Pikachu Lighting", 75000)],
        ),
        (
            "vermilion",
            "bookings.db",
            {"account_id": "A01", "traveler_id": "T07", "after_date": "2026-10-08"},
            [("R100", "A01", "T07", "Ash Ketchum", "2026-10-20")],
        ),
    ):
        sql = (generated / f"keys/{company}.reference.sql").read_text(encoding="utf-8")
        identity_sql, answer_sql = sql.split("-- Answer\n")
        assert not list((generated / "companies" / company).rglob("*.sql"))
        connection = sqlite3.connect(":memory:")
        try:
            connection.execute(
                "ATTACH DATABASE ? AS source",
                (str(generated / "companies" / company / file),),
            )
            if company == "cerulean":
                for table in ("invoices", "invoice_lines", "payments"):
                    connection.execute(
                        f"CREATE TABLE accounting__{table} "
                        f"AS SELECT * FROM source.{table}"
                    )
                workbook_file = "directory.xlsx"
                tab_tables = {
                    "Communities": "directory__communities",
                    "Vendors": "directory__vendors",
                }
                identity_parameters = {"community_name": "Viridian Court"}
                expected_identity = [("H17",)]
            else:
                connection.execute(
                    "CREATE TABLE bookings__reservations AS SELECT * FROM "
                    "source.reservations"
                )
                workbook_file = "clients.xlsx"
                tab_tables = {
                    "Accounts": "clients__accounts",
                    "Travelers": "clients__travelers",
                }
                identity_parameters = {
                    "account_name": "Silph Account",
                    "traveler_name": "Ash Ketchum",
                }
                expected_identity = [("A01", "T07")]
            workbook = load_workbook(
                generated / "companies" / company / workbook_file, read_only=True
            )
            try:
                for tab, table in tab_tables.items():
                    rows = list(workbook[tab].values)
                    columns = ", ".join(f'"{column}" TEXT' for column in rows[0])
                    connection.execute(f"CREATE TABLE {table} ({columns})")
                    placeholders = ", ".join("?" for _ in rows[0])
                    connection.executemany(
                        f"INSERT INTO {table} VALUES ({placeholders})", rows[1:]
                    )
            finally:
                workbook.close()
            assert (
                connection.execute(identity_sql, identity_parameters).fetchall()
                == expected_identity
            )
            assert (
                connection.execute(answer_sql, answer_parameters).fetchall() == expected
            )
        finally:
            connection.close()


def test_every_file_has_reviewed_hash(generated: Path) -> None:
    assert_hash_baseline(DEMO, BASELINE)
    assert set(file_hashes(generated)) == set(read_baseline(BASELINE))


def test_generator_rerun_and_two_fresh_trees(tmp_path: Path) -> None:
    first = tmp_path / "first"
    command = [sys.executable, str(ROOT / "scripts/generate_demo.py"), "--output"]
    subprocess.run([*command, str(first)], cwd=ROOT, check=True, timeout=30)
    before = file_hashes(first)
    baseline_before = BASELINE.read_bytes()
    subprocess.run([*command, str(first)], cwd=ROOT, check=True, timeout=30)
    assert file_hashes(first) == before
    second = tmp_path / "second"
    result = subprocess.run(
        [*command, str(second), "--hashes"],
        cwd=ROOT,
        check=True,
        timeout=30,
        text=True,
        capture_output=True,
    )
    assert file_hashes(second) == before
    assert {
        name: digest
        for digest, name in (line.split("  ", 1) for line in result.stdout.splitlines())
    } == before
    assert BASELINE.read_bytes() == baseline_before


@pytest.mark.parametrize("mutation", ["changed", "added", "removed"])
def test_key_mutation_requires_hash_update(
    generated: Path, tmp_path: Path, mutation: str
) -> None:
    copy = tmp_path / "demo"
    shutil.copytree(DEMO, copy)
    assert_hash_baseline(copy, BASELINE)
    key = copy / "keys/cerulean.answers.json"
    if mutation == "changed":
        keys = read_json(key)
        keys["main"]["rows"][0]["outstanding_cents"] = 1
        key.write_text(json.dumps(keys), encoding="utf-8")
    elif mutation == "added":
        (copy / "keys/unrecorded.json").write_text("{}\n", encoding="utf-8")
    else:
        key.unlink()
    with pytest.raises(AssertionError, match="Hash baseline mismatch"):
        assert_hash_baseline(copy, BASELINE)
    updated = tmp_path / "updated.sha256"
    updated.write_text(
        "".join(f"{digest}  {name}\n" for name, digest in file_hashes(copy).items()),
        encoding="utf-8",
    )
    assert_hash_baseline(copy, updated)
    if mutation == "changed":
        with pytest.raises(AssertionError):
            assert_cerulean_keys(key)


def test_generator_preserves_unowned_files_and_guard_detects_them(
    generated: Path, tmp_path: Path
) -> None:
    copy = tmp_path / "demo"
    shutil.copytree(DEMO, copy)
    assert_hash_baseline(copy, BASELINE)
    extra = copy / "keys/unowned.txt"
    extra.write_text("keep me", encoding="utf-8")
    subprocess.run(
        [sys.executable, str(ROOT / "scripts/generate_demo.py"), "--output", str(copy)],
        cwd=ROOT,
        check=True,
        timeout=30,
    )
    assert extra.read_text(encoding="utf-8") == "keep me"
    with pytest.raises(AssertionError, match="Hash baseline mismatch"):
        assert_hash_baseline(copy, BASELINE)
