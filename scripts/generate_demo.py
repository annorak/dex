#!/usr/bin/env python3
"""Generate the fixed demo sources, hidden keys, and reviewed Pallet package."""

import argparse
import hashlib
import io
import json
import sqlite3
import sys
from collections.abc import Sequence
from datetime import datetime
from pathlib import Path
from typing import NamedTuple
from xml.etree import ElementTree
from zipfile import ZIP_STORED, ZipFile, ZipInfo

import yaml
from openpyxl import Workbook

type Cell = str | int | float
type Row = tuple[Cell, ...]


class Table(NamedTuple):
    name: str
    columns: tuple[str, ...]
    rows: tuple[Row, ...]


FIXED_TIME = datetime(2026, 9, 30)

CERULEAN_COMMUNITIES: Table = Table(
    "Communities",
    ("community_id", "name"),
    (("H17", "Viridian Court"), ("H71", "Viridian Court II"), ("H18", "Pewter Court")),
)
CERULEAN_VENDORS: Table = Table(
    "Vendors",
    ("community_id", "vendor_id", "name"),
    (
        ("H17", "V3", "Bulbasaur Landscaping"),
        ("H17", "V7", "Pikachu Lighting"),
        ("H71", "V7", "Squirtle Pool Care"),
    ),
)
CERULEAN_INVOICES: Table = Table(
    "invoices",
    (
        "invoice_id",
        "community_id",
        "vendor_id",
        "invoice_date",
        "total_cents",
        "status",
    ),
    (
        ("I099", "H17", "V7", "2025-11-03", 15000, "C"),
        ("I100", "H17", "V7", "2026-07-10", 125000, "C"),
        ("I101", "H17", "V7", "2026-08-01", 20000, "C"),
        ("I102", "H17", "V7", "2026-08-15", 10000, "V"),
        ("I103", "H17", "V3", "2026-06-01", 40000, "C"),
        ("I200", "H71", "V7", "2026-08-20", 99900, "C"),
    ),
)
CERULEAN_LINES: Table = Table(
    "invoice_lines",
    ("invoice_id", "line_no", "description", "amount_cents"),
    (
        ("I099", 1, "Invoice total", 15000),
        ("I100", 1, "LED fixtures", 100000),
        ("I100", 2, "Installation labor", 25000),
        ("I101", 1, "Invoice total", 20000),
        ("I102", 1, "Invoice total", 10000),
        ("I103", 1, "Invoice total", 40000),
        ("I200", 1, "Invoice total", 99900),
    ),
)
CERULEAN_PAYMENTS: Table = Table(
    "payments",
    ("payment_id", "invoice_id", "paid_date", "amount_cents"),
    (
        ("PAY1", "I099", "2025-12-01", 15000),
        ("PAY2", "I103", "2026-06-20", 40000),
        ("PAY3", "I100", "2026-07-25", 20000),
        ("PAY4", "I100", "2026-09-05", 30000),
        ("PAY5", "I101", "2026-09-12", 20000),
    ),
)
CERULEAN_OLD: Table = Table(
    "Open Invoices",
    ("Community", "Vendor", "Invoice", "Amount Open"),
    (
        ("Viridian Court", "Pikachu Lighting", "I100", 1050.00),
        ("Viridian Court", "Pikachu Lighting", "I101", 200.00),
        ("Viridian Court II", "Squirtle Pool Care", "I200", 999.00),
    ),
)
VERMILION_ACCOUNTS: Table = Table(
    "Accounts",
    ("account_id", "name"),
    (("A01", "Silph Account"), ("A02", "Devon Account")),
)
VERMILION_TRAVELERS: Table = Table(
    "Travelers",
    ("account_id", "traveler_id", "name"),
    (
        ("A01", "T07", "Ash Ketchum"),
        ("A01", "T09", "Brock"),
        ("A02", "T07", "Ash Ketchum"),
    ),
)
VERMILION_RESERVATIONS: Table = Table(
    "reservations",
    ("reservation_id", "account_id", "traveler_id", "depart_date", "status"),
    (
        ("R100", "A01", "T07", "2026-10-20", "C"),
        ("R101", "A01", "T07", "2026-10-22", "X"),
        ("R102", "A01", "T07", "2026-09-25", "C"),
        ("R103", "A01", "T07", "2026-10-08", "C"),
        ("R104", "A01", "T07", "2026-11-03", "X"),
        ("R200", "A02", "T07", "2026-10-20", "C"),
        ("R201", "A02", "T07", "2026-10-15", "X"),
        ("R300", "A01", "T07", "2026-10-22", "P"),
    ),
)
VERMILION_OLD: Table = Table(
    "Trips",
    ("Traveler", "Account", "Reservation", "Depart", "Status"),
    (
        ("Ash Ketchum", "Silph Account", "R100", "2026-10-20", "Confirmed"),
        ("Ash Ketchum", "Silph Account", "R101", "2026-10-22", "Confirmed"),
    ),
)
PALLET_COMMUNITIES: Table = Table(
    "communities",
    ("community_id", "name"),
    (("P01", "Pallet Grove"), ("P02", "Lavender Grove")),
)
PALLET_VENDORS: Table = Table(
    "vendors",
    ("vendor_id", "name"),
    (("V1", "Eevee Maintenance"),),
)
PALLET_INVOICES: Table = Table(
    "invoices",
    (
        "invoice_id",
        "community_id",
        "vendor_id",
        "invoice_date",
        "total_cents",
        "balance_cents",
        "status",
    ),
    (
        ("P100", "P01", "V1", "2026-09-01", 20000, 15000, "O"),
        ("P101", "P01", "V1", "2026-08-01", 10000, 0, "C"),
    ),
)

CERULEAN_DOC = (
    "1. Vendor bills are entered as invoices, and each payment "
    "is recorded as its own entry against an invoice.\n2. All "
    "amounts in the accounting system are in cents.\n3. Each "
    'month, the AP clerk exports an "Open Invoices" sheet for '
    "the board.\n"
)
VERMILION_DOC = (
    "1. Cancelled bookings are marked `X`.\n2. Each client "
    "company has its own account, and traveler numbers start "
    "over in each account.\n"
)
PALLET_DOC = (
    "1. Invoices store the amount still owed in balance_cents. "
    "Amounts are whole US cents.\n2. Invoice status O means "
    "open. C means closed and paid in full.\n3. Vendor IDs are "
    "unique across Pallet HOA. The accounting database is the "
    "trusted source.\n"
)
CERULEAN_ANSWERS = (
    (
        "status_meaning",
        "accounting.invoices.status",
        (
            "C means current, so the invoice is posted and counts. V "
            "means void, so ignore it."
        ),
    ),
    (
        "balance_rule",
        "accounting.invoices",
        (
            "There's no balance field. What's owed is the total minus "
            "the payments recorded against the invoice. An invoice is "
            "unpaid if it isn't void and something is still owed."
        ),
    ),
    (
        "trusted_source",
        "open_invoices_aug",
        (
            "Use the accounting system for anything about balances. The"
            " Open Invoices sheet is a monthly export for the board, so"
            " it goes stale."
        ),
    ),
    (
        "id_scope",
        "directory.vendors.vendor_id",
        (
            "Vendor numbers are assigned per community. V7 at one "
            "community is a different vendor from V7 at another."
        ),
    ),
    (
        "current_process",
        "unpaid_invoices_by_community",
        (
            "Board members ask before their meetings. The AP clerk "
            "emails them the latest Open Invoices sheet."
        ),
    ),
)
VERMILION_ANSWERS = (
    (
        "status_meaning",
        "bookings.reservations.status",
        (
            "C means confirmed (booked). P means pending: on hold, not "
            "booked yet. X means cancelled."
        ),
    ),
    (
        "trusted_source",
        "trips_old",
        "Use the bookings system. The Trips sheet stopped being updated in September.",
    ),
    (
        "date_rule",
        "bookings.reservations.depart_date",
        "It's the day the trip starts. We don't track times here.",
    ),
    (
        "current_process",
        "upcoming_reservations",
        "Account managers check the bookings system before a client call.",
    ),
)
PALLET_ANSWERS = (
    (
        "status_meaning",
        "accounting.invoices.status",
        "O means open. C means closed and paid in full.",
    ),
    (
        "balance_rule",
        "accounting.invoices.balance_cents",
        (
            "Use the stored balance_cents. A positive balance on an "
            "open invoice is unpaid."
        ),
    ),
    (
        "id_scope",
        "accounting.vendors.vendor_id",
        "Vendor IDs are unique across Pallet HOA.",
    ),
    ("trusted_source", "accounting", "Use the accounting database for balances."),
)
HOA_CHECKLIST = {
    "industry": "hoa",
    "lookup": "unpaid_invoices_by_community",
    "version": 1,
    "question": (
        "Which invoices are still unpaid for a community, and which"
        " vendors are they for?"
    ),
    "inputs": {"community_name": "the exact name of one community"},
    "outputs": [
        "invoice_id",
        "vendor_id",
        "vendor_name",
        "outstanding_cents",
        "currency",
    ],
    "ask_employee": [
        "What does each invoice status code mean? Codes differ between companies.",
        "Is a balance stored on each invoice, or is it worked out from payments?",
        "If two sources disagree about what's owed, which one do you trust?",
    ],
    "watch_for": [
        "Community names can look alike. Match the exact name, then confirm the ID.",
        "Check whether amounts are stored in cents or dollars.",
    ],
    "checks": ["identity_is_unique", "exact_rows_and_values", "known_empty_is_ok"],
}
CERULEAN_SCHEMA = (
    "CREATE TABLE invoices (invoice_id TEXT PRIMARY KEY, "
    "community_id TEXT NOT NULL, vendor_id TEXT NOT NULL, "
    "invoice_date TEXT NOT NULL, total_cents INTEGER NOT NULL, "
    "status TEXT NOT NULL);\nCREATE TABLE invoice_lines "
    "(invoice_id TEXT NOT NULL REFERENCES invoices(invoice_id),"
    " line_no INTEGER NOT NULL, description TEXT NOT NULL, "
    "amount_cents INTEGER NOT NULL, PRIMARY KEY (invoice_id, "
    "line_no));\nCREATE TABLE payments (payment_id TEXT PRIMARY "
    "KEY, invoice_id TEXT NOT NULL REFERENCES "
    "invoices(invoice_id), paid_date TEXT NOT NULL, "
    "amount_cents INTEGER NOT NULL);\n"
)
VERMILION_SCHEMA = (
    "CREATE TABLE reservations (reservation_id TEXT PRIMARY "
    "KEY, account_id TEXT NOT NULL, traveler_id TEXT NOT NULL, "
    "depart_date TEXT NOT NULL, status TEXT NOT NULL);\n"
)
PALLET_SCHEMA = (
    "CREATE TABLE communities (community_id TEXT PRIMARY KEY, "
    "name TEXT NOT NULL);\nCREATE TABLE vendors (vendor_id TEXT "
    "PRIMARY KEY, name TEXT NOT NULL);\nCREATE TABLE invoices "
    "(invoice_id TEXT PRIMARY KEY, community_id TEXT NOT NULL "
    "REFERENCES communities(community_id), vendor_id TEXT NOT "
    "NULL REFERENCES vendors(vendor_id), invoice_date TEXT NOT "
    "NULL, total_cents INTEGER NOT NULL, balance_cents INTEGER "
    "NOT NULL, status TEXT NOT NULL);\n"
)
PALLET_IDENTITY = """SELECT community_id FROM communities WHERE name = :community_name;
"""
PALLET_QUERY = (
    "SELECT i.invoice_id, i.vendor_id, v.name AS vendor_name,"
    "\n       i.balance_cents AS outstanding_cents, 'USD' AS "
    "currency\nFROM invoices AS i\nJOIN vendors AS v ON "
    "v.vendor_id = i.vendor_id\nWHERE i.community_id = "
    ":community_id AND i.status = 'O' AND i.balance_cents > 0\n"
    "ORDER BY i.invoice_id;\n"
)
CERULEAN_REFERENCE = (
    "-- Identity\nSELECT community_id FROM "
    "directory__communities WHERE name = :community_name;\n\n-- "
    "Answer\nWITH paid AS (\n  SELECT invoice_id, "
    "SUM(amount_cents) AS paid_cents\n  FROM "
    "accounting__payments GROUP BY invoice_id\n)\nSELECT "
    "i.invoice_id, i.vendor_id, v.name AS vendor_name,\n       "
    "i.total_cents - COALESCE(paid.paid_cents, 0) AS "
    "outstanding_cents\nFROM accounting__invoices AS i\nLEFT JOIN"
    " paid ON paid.invoice_id = i.invoice_id\nJOIN "
    "directory__vendors AS v\n  ON v.community_id = "
    "i.community_id AND v.vendor_id = i.vendor_id\nWHERE "
    "i.community_id = :community_id AND i.status <> 'V'\n  AND "
    "i.total_cents - COALESCE(paid.paid_cents, 0) > 0\nORDER BY "
    "i.invoice_id;\n"
)
VERMILION_REFERENCE = (
    "-- Identity\nSELECT t.account_id, t.traveler_id\nFROM "
    "clients__travelers AS t\nJOIN clients__accounts AS a ON "
    "a.account_id = t.account_id\nWHERE a.name = :account_name "
    "AND t.name = :traveler_name;\n\n-- Answer\nSELECT "
    "r.reservation_id, r.account_id, r.traveler_id,\n       "
    "t.name AS traveler_name, r.depart_date\nFROM "
    "bookings__reservations AS r\nJOIN clients__travelers AS t"
    "\n  ON t.account_id = r.account_id AND t.traveler_id = "
    "r.traveler_id\nWHERE r.account_id = :account_id AND "
    "r.traveler_id = :traveler_id\n  AND r.status = 'C' AND "
    "r.depart_date > :after_date\nORDER BY r.depart_date, "
    "r.reservation_id;\n"
)


def write_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(value, encoding="utf-8")


def write_json(path: Path, value: object) -> None:
    write_text(
        path, json.dumps(value, indent=2, sort_keys=True, ensure_ascii=False) + "\n"
    )


def write_yaml(path: Path, value: object) -> None:
    write_text(
        path, yaml.safe_dump(value, sort_keys=False, allow_unicode=True, width=1000)
    )


def write_database(path: Path, schema: str, tables: tuple[Table, ...]) -> None:
    connection = sqlite3.connect(":memory:")
    try:
        connection.execute("PRAGMA foreign_keys = ON")
        connection.executescript(schema)
        for name, columns, rows in tables:
            placeholders = ", ".join("?" for _ in columns)
            connection.executemany(
                f'INSERT INTO "{name}" VALUES ({placeholders})', rows
            )
        connection.commit()
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(connection.serialize())
    finally:
        connection.close()


def write_workbook(path: Path, sheets: tuple[Table, ...]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    for name, columns, rows in sheets:
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        for cells in sheet.columns:
            sheet.column_dimensions[cells[0].column_letter].width = min(
                45, max(len(str(cell.value)) for cell in cells) + 2
            )
        if "Amount Open" in columns:
            for row in sheet.iter_rows(min_row=2, min_col=4, max_col=4):
                row[0].number_format = "0.00"
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    path.parent.mkdir(parents=True, exist_ok=True)
    with (
        ZipFile(buffer) as original,
        ZipFile(path, "w", compression=ZIP_STORED) as canonical,
    ):
        for name in sorted(original.namelist()):
            contents = original.read(name)
            if name == "docProps/core.xml":
                # openpyxl sets modified to the wall clock during save.
                ElementTree.register_namespace(
                    "cp",
                    (
                        "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
                    ),
                )
                ElementTree.register_namespace("dc", "http://purl.org/dc/elements/1.1/")
                ElementTree.register_namespace("dcterms", "http://purl.org/dc/terms/")
                ElementTree.register_namespace(
                    "xsi", "http://www.w3.org/2001/XMLSchema-instance"
                )
                root = ElementTree.fromstring(contents)
                modified = root.find("{http://purl.org/dc/terms/}modified")
                assert modified is not None
                modified.text = FIXED_TIME.isoformat() + "Z"
                contents = ElementTree.tostring(root, encoding="utf-8")
            member = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            member.create_system = 3
            member.external_attr = 0o100644 << 16
            canonical.writestr(member, contents)


def source_ref(
    company: str, scope: str, source: str, table: str, record: str
) -> dict[str, str]:
    return {
        "source": source,
        "table": table,
        "record_id": f"{company}/{scope}/{record}",
    }


def invoice_answer(
    scope: str,
    invoice: str,
    vendor: str,
    name: str,
    amount: int,
    payments: tuple[str, ...],
) -> dict[str, object]:
    return {
        "invoice_id": f"cerulean/{scope}/{invoice}",
        "vendor_id": f"cerulean/{scope}/{vendor}",
        "vendor_name": name,
        "outstanding_cents": amount,
        "currency": "USD",
        "source_refs": [
            source_ref("cerulean", scope, "accounting", "invoices", invoice),
            *[
                source_ref("cerulean", scope, "accounting", "payments", payment)
                for payment in payments
            ],
        ],
    }


def reservation_answer(scope: str, reservation: str, account: str) -> dict[str, object]:
    return {
        "reservation_id": f"vermilion/{scope}/{reservation}",
        "traveler_id": f"vermilion/{scope}/T07",
        "traveler_name": "Ash Ketchum",
        "account_id": f"vermilion/{scope}",
        "account_name": account,
        "depart_date": "2026-10-20",
        "source_refs": [
            source_ref("vermilion", scope, "bookings", "reservations", reservation)
        ],
    }


def case(
    name: str,
    inputs: dict[str, str],
    status: str,
    identity: list[str],
    rows: list[dict[str, object]],
) -> dict[str, object]:
    return {
        "check": name,
        "input": inputs,
        "status": status,
        "identity": identity,
        "rows": rows,
    }


def write_company_sources(output: Path) -> None:
    cerulean = output / "companies/cerulean"
    write_database(
        cerulean / "accounting.db",
        CERULEAN_SCHEMA,
        (CERULEAN_INVOICES, CERULEAN_LINES, CERULEAN_PAYMENTS),
    )
    write_workbook(
        cerulean / "directory.xlsx", (CERULEAN_COMMUNITIES, CERULEAN_VENDORS)
    )
    write_workbook(cerulean / "Open Invoices - Aug 31 2026.xlsx", (CERULEAN_OLD,))
    write_text(cerulean / "docs/ap_process.md", CERULEAN_DOC)
    vermilion = output / "companies/vermilion"
    write_database(
        vermilion / "bookings.db", VERMILION_SCHEMA, (VERMILION_RESERVATIONS,)
    )
    write_workbook(
        vermilion / "clients.xlsx", (VERMILION_ACCOUNTS, VERMILION_TRAVELERS)
    )
    write_workbook(vermilion / "Trips (old).xlsx", (VERMILION_OLD,))
    write_text(vermilion / "docs/booking_process.md", VERMILION_DOC)
    pallet = output / "companies/pallet"
    write_database(
        pallet / "accounting.db",
        PALLET_SCHEMA,
        (PALLET_COMMUNITIES, PALLET_VENDORS, PALLET_INVOICES),
    )
    write_text(pallet / "docs/ap_process.md", PALLET_DOC)
    for company, date, sources in (
        (
            "cerulean",
            "2026-09-30",
            [
                {
                    "id": "accounting",
                    "file": "accounting.db",
                    "tables": ["invoices", "invoice_lines", "payments"],
                },
                {
                    "id": "directory",
                    "file": "directory.xlsx",
                    "tabs": ["Communities", "Vendors"],
                },
                {
                    "id": "open_invoices_aug",
                    "file": "Open Invoices - Aug 31 2026.xlsx",
                    "tabs": ["Open Invoices"],
                },
                {"id": "docs", "file": "docs/ap_process.md"},
            ],
        ),
        (
            "vermilion",
            "2026-10-08",
            [
                {"id": "bookings", "file": "bookings.db", "tables": ["reservations"]},
                {
                    "id": "clients",
                    "file": "clients.xlsx",
                    "tabs": ["Accounts", "Travelers"],
                },
                {"id": "trips_old", "file": "Trips (old).xlsx", "tabs": ["Trips"]},
                {"id": "docs", "file": "docs/booking_process.md"},
            ],
        ),
        (
            "pallet",
            "2026-09-30",
            [
                {
                    "id": "accounting",
                    "file": "accounting.db",
                    "tables": ["communities", "vendors", "invoices"],
                },
                {"id": "docs", "file": "docs/ap_process.md"},
            ],
        ),
    ):
        write_json(
            output / f"companies/{company}/sources.json",
            {
                "company": company,
                "snapshot": f"{company}-v1",
                "as_of": date,
                "sources": sources,
            },
        )


def write_keys(output: Path) -> None:
    main = invoice_answer(
        "H17", "I100", "V7", "Pikachu Lighting", 75000, ("PAY3", "PAY4")
    )
    lookalike = invoice_answer("H71", "I200", "V7", "Squirtle Pool Care", 99900, ())
    cerulean_cases = [
        case(
            check, {"community_name": "Viridian Court"}, "ok", ["cerulean/H17"], [main]
        )
        for check in ("identity_viridian", "rows_viridian", "values_viridian")
    ]
    cerulean_cases += [
        case(
            "lookalike_viridian_ii",
            {"community_name": "Viridian Court II"},
            "ok",
            ["cerulean/H71"],
            [lookalike],
        ),
        case(
            "empty_pewter",
            {"community_name": "Pewter Court"},
            "ok",
            ["cerulean/H18"],
            [],
        ),
        case(
            "unknown_name",
            {"community_name": "Viridian Ct"},
            "unresolved_identity",
            [],
            [],
        ),
        case(
            "source_refs",
            {"community_name": "Viridian Court"},
            "ok",
            ["cerulean/H17"],
            [main],
        ),
    ]
    write_json(
        output / "keys/cerulean.answers.json",
        {
            "company": "cerulean",
            "snapshot": "cerulean-v1",
            "as_of": "2026-09-30",
            "lookup": "unpaid_invoices_by_community",
            "main": cerulean_cases[1],
            "checks": cerulean_cases,
            "follow_up": {
                "question": (
                    "How much has each Viridian Court vendor been paid in 2026,"
                    " counting payments dated January 1 through September 30?"
                ),
                "input": {
                    "community_name": "Viridian Court",
                    "from_date": "2026-01-01",
                    "through_date": "2026-09-30",
                },
                "status": "ok",
                "rows": [
                    {
                        "vendor_id": "cerulean/H17/V3",
                        "vendor_name": "Bulbasaur Landscaping",
                        "paid_cents": 40000,
                        "currency": "USD",
                    },
                    {
                        "vendor_id": "cerulean/H17/V7",
                        "vendor_name": "Pikachu Lighting",
                        "paid_cents": 70000,
                        "currency": "USD",
                    },
                ],
            },
        },
    )
    travel_main = reservation_answer("A01", "R100", "Silph Account")
    travel_other = reservation_answer("A02", "R200", "Devon Account")
    inputs = {
        "traveler_name": "Ash Ketchum",
        "account_name": "Silph Account",
        "after_date": "2026-10-08",
    }
    travel_cases = [
        case(check, inputs, "ok", ["vermilion/A01/T07"], [travel_main])
        for check in ("identity_ash_silph", "rows_ash_silph", "values_ash_silph")
    ]
    travel_cases += [
        case(
            "same_name_devon",
            {**inputs, "account_name": "Devon Account"},
            "ok",
            ["vermilion/A02/T07"],
            [travel_other],
        ),
        case(
            "empty_brock",
            {**inputs, "traveler_name": "Brock"},
            "ok",
            ["vermilion/A01/T09"],
            [],
        ),
        case(
            "unknown_traveler",
            {**inputs, "traveler_name": "Ash K."},
            "unresolved_identity",
            [],
            [],
        ),
        case("source_refs", inputs, "ok", ["vermilion/A01/T07"], [travel_main]),
    ]
    write_json(
        output / "keys/vermilion.answers.json",
        {
            "company": "vermilion",
            "snapshot": "vermilion-v1",
            "as_of": "2026-10-08",
            "lookup": "upcoming_reservations",
            "main": travel_cases[1],
            "checks": travel_cases,
            "follow_up": {
                "question": (
                    "Which Silph Account travelers have a cancelled reservation"
                    " that was scheduled to depart in October 2026?"
                ),
                "input": {
                    "account_name": "Silph Account",
                    "from_date": "2026-10-01",
                    "through_date": "2026-10-31",
                },
                "status": "ok",
                "rows": [
                    {
                        "traveler_id": "vermilion/A01/T07",
                        "traveler_name": "Ash Ketchum",
                        "account_id": "vermilion/A01",
                        "account_name": "Silph Account",
                    }
                ],
            },
        },
    )
    pallet_row: dict[str, object] = {
        "invoice_id": "pallet/P01/P100",
        "vendor_id": "pallet/P01/V1",
        "vendor_name": "Eevee Maintenance",
        "outstanding_cents": 15000,
        "currency": "USD",
        "source_refs": [source_ref("pallet", "P01", "accounting", "invoices", "P100")],
    }
    write_json(
        output / "keys/pallet.answers.json",
        {
            "company": "pallet",
            "snapshot": "pallet-v1",
            "as_of": "2026-09-30",
            "lookup": "unpaid_invoices_by_community",
            "main": case(
                "rows_pallet",
                {"community_name": "Pallet Grove"},
                "ok",
                ["pallet/P01"],
                [pallet_row],
            ),
            "empty": case(
                "empty_lavender",
                {"community_name": "Lavender Grove"},
                "ok",
                ["pallet/P02"],
                [],
            ),
            "unknown": case(
                "unknown_name",
                {"community_name": "Pallet Gr"},
                "unresolved_identity",
                [],
                [],
            ),
        },
    )
    for company, answers in (
        ("cerulean", CERULEAN_ANSWERS),
        ("vermilion", VERMILION_ANSWERS),
        ("pallet", PALLET_ANSWERS),
    ):
        write_yaml(
            output / f"keys/{company}.employee.yaml",
            {
                "company": company,
                "answers": [
                    {"topic": topic, "target": target, "answer": answer}
                    for topic, target, answer in answers
                ],
            },
        )
    write_text(output / "keys/cerulean.reference.sql", CERULEAN_REFERENCE)
    write_text(output / "keys/vermilion.reference.sql", VERMILION_REFERENCE)


def write_pallet_package(output: Path) -> None:
    package = output / "packages/pallet"
    source = output / "companies/pallet"
    for original, target in (
        ("accounting.db", "accounting.db"),
        ("docs/ap_process.md", "ap_process.md"),
    ):
        path = package / "evidence" / target
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes((source / original).read_bytes())
    write_text(package / "queries/identity.sql", PALLET_IDENTITY)
    write_text(package / "queries/answer.sql", PALLET_QUERY)
    write_json(
        package / "queries/settings.json",
        {"database": "evidence/accounting.db", "read_only": True, "timeout_seconds": 2},
    )
    write_text(
        package / "overview.md",
        (
            "# Pallet HOA reviewed example\n\nPrepared before the demo. "
            "This is an invented, simulated review, not a live "
            "onboarding result.\n\nSnapshot pallet-v1 is as of "
            "2026-09-30. The accounting database stores communities, "
            "company-wide unique vendors, and invoices.\nUse "
            "invoices.balance_cents, in whole US cents, for the amount "
            "owed. O means open and C means closed, paid in full.\nThese"
            " rules are documented in evidence/ap_process.md items 1-3 "
            "and confirmed by employee review pallet-employee-1.\nJoin "
            "invoices to vendors on vendor_id. Resolve a community by "
            "exact name. See queries/identity.sql and "
            "queries/answer.sql.\nSupported lookup: "
            "unpaid_invoices_by_community. There are no open questions "
            "for this lookup.\nThe engineer approved the saved queries "
            "and three passing checks in pallet-engineer-1. Results are"
            " in checks/results.json.\nLimits: fixed invented snapshot "
            "only; no historical balances, live permissions, or live "
            "freshness guarantees.\n"
        ),
    )
    write_json(
        package / "sources.json",
        {
            "company": "pallet",
            "snapshot": "pallet-v1",
            "as_of": "2026-09-30",
            "sources": [
                {
                    "id": "accounting",
                    "original_file": "accounting.db",
                    "file": "evidence/accounting.db",
                    "tables": {
                        "communities": {
                            "primary_key": ["community_id"],
                            "records": ["pallet/P01", "pallet/P02"],
                        },
                        "vendors": {
                            "primary_key": ["vendor_id"],
                            "records": ["pallet/P01/V1"],
                        },
                        "invoices": {
                            "primary_key": ["invoice_id"],
                            "records": ["pallet/P01/P100", "pallet/P01/P101"],
                        },
                    },
                },
                {
                    "id": "docs",
                    "original_file": "docs/ap_process.md",
                    "file": "evidence/ap_process.md",
                    "items": [1, 2, 3],
                },
            ],
        },
    )
    write_json(
        package / "lookups.json",
        {
            "lookups": [
                {
                    "id": "unpaid_invoices_by_community",
                    "question": HOA_CHECKLIST["question"],
                    "inputs": HOA_CHECKLIST["inputs"],
                    "outputs": HOA_CHECKLIST["outputs"],
                    "identity_query": "queries/identity.sql",
                    "answer_query": "queries/answer.sql",
                    "rules": [
                        "Exact community name must resolve once.",
                        "O is open; C is closed.",
                        "Positive stored balance_cents on an open invoice is unpaid.",
                        "Vendor IDs are unique across the company.",
                        "Amounts are whole US cents; currency is USD.",
                    ],
                }
            ]
        },
    )
    write_json(
        package / "decisions.json",
        {
            "decisions": [
                {
                    "id": "pallet-employee-1",
                    "role": "employee",
                    "simulated": True,
                    "outcome": "approved",
                    "reviewed": [
                        "status codes",
                        "stored balance",
                        "vendor scope",
                        "trusted source",
                    ],
                    "at": "2026-09-30T00:00:00Z",
                },
                {
                    "id": "pallet-engineer-1",
                    "role": "engineer",
                    "simulated": True,
                    "outcome": "approved",
                    "reviewed": [
                        "queries/identity.sql",
                        "queries/answer.sql",
                        "checks/results.json",
                    ],
                    "at": "2026-09-30T00:00:00Z",
                },
            ]
        },
    )
    checks = [
        {
            "id": "identity_is_unique",
            "community_name": "Pallet Grove",
            "expected_identity": ["P01"],
            "expected_rows": [["P100", "V1", "Eevee Maintenance", 15000, "USD"]],
        },
        {
            "id": "exact_rows_and_values",
            "community_name": "Pallet Grove",
            "expected_identity": ["P01"],
            "expected_rows": [["P100", "V1", "Eevee Maintenance", 15000, "USD"]],
        },
        {
            "id": "known_empty_is_ok",
            "community_name": "Lavender Grove",
            "expected_identity": ["P02"],
            "expected_rows": [],
        },
    ]
    write_json(package / "checks/definitions.json", {"checks": checks})
    results: list[dict[str, object]] = []
    connection = sqlite3.connect(
        f"{(package / 'evidence/accounting.db').as_uri()}?mode=ro", uri=True
    )
    try:
        for check in checks:
            identities = connection.execute(
                PALLET_IDENTITY, {"community_name": check["community_name"]}
            ).fetchall()
            assert len(identities) == 1
            rows = connection.execute(
                PALLET_QUERY, {"community_id": identities[0][0]}
            ).fetchall()
            identity = [row[0] for row in identities]
            actual = [list(row) for row in rows]
            assert (
                identity == check["expected_identity"]
                and actual == check["expected_rows"]
            )
            results.append(
                {
                    "id": check["id"],
                    "status": "passed",
                    "actual_identity": identity,
                    "actual_rows": actual,
                }
            )
    finally:
        connection.close()
    write_json(
        package / "checks/results.json",
        {"run": "pallet-reviewed-v1", "results": results},
    )
    hashes = {
        path.relative_to(package).as_posix(): hashlib.sha256(
            path.read_bytes()
        ).hexdigest()
        for path in sorted(package.rglob("*"))
        if path.is_file() and path.name != "manifest.json"
    }
    write_json(
        package / "manifest.json",
        {
            "package_format": 1,
            "dex_version": "0.1.0",
            "company": "pallet",
            "question": HOA_CHECKLIST["question"],
            "lookup": "unpaid_invoices_by_community",
            "snapshot": "pallet-v1",
            "as_of": "2026-09-30",
            "query_version": 1,
            "checks_version": 1,
            "checklist_used": {
                "industry": "hoa",
                "lookup": "unpaid_invoices_by_community",
                "version": 1,
            },
            "prepared_example": True,
            "simulated": True,
            "reviews": ["pallet-employee-1", "pallet-engineer-1"],
            "limits": [
                "Fixed invented snapshot",
                "No live permissions or freshness validation",
                "No historical balances",
            ],
            "files": hashes,
        },
    )


def generate_demo(output: Path) -> None:
    write_company_sources(output)
    write_keys(output)
    write_yaml(output / "checklists/hoa/unpaid_invoices.v1.yaml", HOA_CHECKLIST)
    write_pallet_package(output)


def main(argv: Sequence[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("demo"))
    parser.add_argument(
        "--hashes",
        action="store_true",
        help="Print every generated tree file hash; never update the reviewed baseline",
    )
    args = parser.parse_args(argv)
    args.output = args.output.resolve()
    generate_demo(args.output)
    if args.hashes:
        for path in sorted(args.output.rglob("*")):
            if path.is_file():
                digest = hashlib.sha256(path.read_bytes()).hexdigest()
                name = path.relative_to(args.output).as_posix()
                print(f"{digest}  {name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
